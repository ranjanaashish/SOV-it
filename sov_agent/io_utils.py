"""File ingestion (FR-1). Reads every sheet WITHOUT assuming a header row.

Returns raw grids (pandas DataFrame, object dtype, 0-based positional index)
plus merged-cell ranges per sheet so Agent 1 can reason about layout.
"""
from __future__ import annotations

import io
from pathlib import Path

import pandas as pd

MAX_ROWS = 200_000          # hard safety cap; FR-1 needs 5,000
MAX_BYTES = 50 * 1024 * 1024


class IngestError(ValueError):
    """User-readable ingestion failure (NFR-4)."""


def read_any(data: bytes, file_name: str) -> tuple[dict[str, pd.DataFrame], dict[str, list[str]]]:
    if not data:
        raise IngestError("The uploaded file is empty.")
    if len(data) > MAX_BYTES:
        raise IngestError(f"File is larger than {MAX_BYTES // (1024 * 1024)} MB.")
    ext = Path(file_name).suffix.lower()
    is_zip = data[:2] == b"PK"                        # xlsx/xlsm are zip containers
    if ext in (".xlsx", ".xlsm") or (is_zip and ext not in (".csv", ".txt")):
        if not is_zip:
            raise IngestError(f"'{file_name}' has an Excel extension but is not a valid .xlsx file.")
        return _read_xlsx(data)
    if ext in (".csv", ".txt"):
        if is_zip:
            raise IngestError(f"'{file_name}' looks like an Excel workbook saved with a .csv extension. Rename it to .xlsx.")
        return _read_csv(data, file_name)
    if ext == ".xls":
        try:
            sheets = pd.read_excel(io.BytesIO(data), sheet_name=None, header=None, dtype=object)
        except Exception as exc:
            raise IngestError("Legacy .xls needs the 'xlrd' package; please save the file as .xlsx.") from exc
        return {k: _trim(v) for k, v in sheets.items()}, {k: [] for k in sheets}
    raise IngestError(f"Unsupported file type '{ext or 'unknown'}'. Please upload .xlsx or .csv.")


def _read_xlsx(data: bytes):
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=False)
    except Exception as exc:
        raise IngestError(f"Could not open the workbook: {exc}") from exc
    sheets, merges = {}, {}
    for ws in wb.worksheets:
        rows = []
        for i, row in enumerate(ws.iter_rows(min_row=1, min_col=1, values_only=True)):
            if i >= MAX_ROWS:
                break
            rows.append(list(row))
        df = pd.DataFrame(rows, dtype=object) if rows else pd.DataFrame(dtype=object)
        sheets[ws.title] = _trim(df)
        merges[ws.title] = [str(r) for r in ws.merged_cells.ranges]
    if not sheets:
        raise IngestError("The workbook contains no worksheets.")
    return sheets, merges


def _read_csv(data: bytes, file_name: str):
    last = None
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            df = pd.read_csv(io.BytesIO(data), header=None, dtype=object, encoding=enc,
                             sep=None, engine="python", skip_blank_lines=False, nrows=MAX_ROWS)
            name = Path(file_name).stem or "CSV"
            return {name: _trim(df)}, {name: []}
        except Exception as exc:
            last = exc
    raise IngestError(f"Could not parse the CSV file: {last}")


def _trim(df: pd.DataFrame) -> pd.DataFrame:
    """Normalise blanks to None and cut trailing empty rows/cols.
    Positions are preserved so merged-cell coordinates stay valid."""
    if df.empty:
        return df
    df = df.map(lambda v: None if (v is None or (isinstance(v, float) and pd.isna(v))
                                    or (isinstance(v, str) and not v.strip())) else v)
    nonempty = df.notna()
    if not nonempty.values.any():
        return pd.DataFrame(dtype=object)
    last_row = int(nonempty.any(axis=1).values.nonzero()[0].max())
    last_col = int(nonempty.any(axis=0).values.nonzero()[0].max())
    out = df.iloc[: last_row + 1, : last_col + 1].copy()
    out.index = range(out.shape[0])
    out.columns = range(out.shape[1])
    return out


def merged_spans(merged: list[str], row0: int) -> list[tuple[int, int]]:
    """Column spans (0-based, inclusive) of merged ranges that cover sheet row `row0`."""
    from openpyxl.utils.cell import range_boundaries

    spans = []
    for ref in merged:
        min_col, min_row, max_col, max_row = range_boundaries(ref)
        if min_row - 1 <= row0 <= max_row - 1 and max_col > min_col:
            spans.append((min_col - 1, max_col - 1))
    return spans