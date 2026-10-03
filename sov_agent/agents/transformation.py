"""Agent 4 - Controlled Transformation. Plain code, no LLM calls.

* Runs only when the review gate is validated (no pending / escalated items).
* Works on a copy; the source DataFrame is never mutated.
* Applies approved / edited recommendations only, logs every change.
* Reindexes to the 17 fields in fixed order, casts types, keeps blanks blank.
"""
from __future__ import annotations

import io
import json
import numbers
from typing import Any

import pandas as pd
from openpyxl import Workbook, load_workbook

from .. import schema
from ..state import AuditEntry, SOVState, now_iso
from ..text import is_blank, parse_number, to_display
from .data_quality import DROP

NAME = "Agent 4 · Controlled Transformation"
APPLIED = ("approved", "edited")


def _same(cur: Any, before: Any) -> bool:
    if is_blank(cur) and is_blank(before):
        return True
    return cur == before or str(to_display(cur)).strip() == str(to_display(before)).strip()


class TransformationAgent:
    name = NAME

    def run(self, state: SOVState) -> SOVState:
        state.require("quality", "review")
        state.start("transform", NAME)
        src: pd.DataFrame = state.source_df.copy(deep=True)
        audit: list[AuditEntry] = []
        recs = [r for r in state.active_recs() if r.status in APPLIED]

        # 1. column mappings (rename)
        out = pd.DataFrame(index=src.index, columns=schema.TARGET_FIELDS, dtype=object)
        target_src: dict[str, str] = {}
        for r in recs:
            if r.action_type != "column_mapping" or not r.source_column:
                continue
            if r.proposed_target is None:
                audit.append(AuditEntry(source_column=r.source_column, target_column=None,
                                        transformation_applied="exclude_column", confidence=r.confidence,
                                        approved_by=r.decided_by, rationale=r.rationale, recommendation_id=r.id))
                continue
            if r.proposed_target in target_src:
                raise RuntimeError(f"Two approved columns map to {r.proposed_target}")
            target_src[r.proposed_target] = r.source_column
            out[r.proposed_target] = src[r.source_column].astype(object)
            audit.append(AuditEntry(source_column=r.source_column, target_column=r.proposed_target,
                                    transformation_applied="rename_column", before_value=r.source_column,
                                    after_value=r.proposed_target, confidence=r.confidence,
                                    approved_by=r.decided_by, rationale=r.rationale, recommendation_id=r.id))

        # 2. row-level corrections (standardisation after corrections; whitespace last)
        order = {"data_correction": 0, "flag_for_review": 1, "standardisation": 2}
        cell_recs = sorted([r for r in recs if r.action_type != "column_mapping" and r.changes],
                           key=lambda r: (order[r.action_type], r.issue_type == "whitespace"))
        drop_rows: set[int] = set()
        for r in cell_recs:
            for c in r.changes:
                if c.after == DROP:
                    drop_rows.add(c.row)
                    audit.append(AuditEntry(source_column=None, target_column=None, row=c.row,
                                            transformation_applied="exclude_row", before_value="(row)",
                                            after_value="(excluded)", confidence=r.confidence,
                                            approved_by=r.decided_by, rationale=r.rationale, recommendation_id=r.id))
                    continue
                field = r.field
                if field not in out.columns or c.row not in out.index or field not in target_src:
                    continue
                cur = out.at[c.row, field]
                if not _same(cur, c.before):
                    audit.append(AuditEntry(source_column=r.source_column, target_column=field, row=c.row,
                                            transformation_applied=f"skipped_{r.issue_type} (value already changed)",
                                            before_value=to_display(cur), after_value=to_display(cur),
                                            confidence=r.confidence, approved_by=r.decided_by,
                                            rationale="Cell was modified by an earlier approved change.",
                                            recommendation_id=r.id))
                    continue
                out.at[c.row, field] = c.after
                audit.append(AuditEntry(source_column=r.source_column, target_column=field, row=c.row,
                                        transformation_applied=r.issue_type or r.action_type,
                                        before_value=to_display(c.before), after_value=to_display(c.after),
                                        confidence=r.confidence, approved_by=r.decided_by,
                                        rationale=r.rationale, recommendation_id=r.id))
        if drop_rows:
            out = out.drop(index=[r for r in drop_rows if r in out.index])

        # 3. type casting per data dictionary; un-castable values blanked and logged
        for field in schema.TARGET_FIELDS:
            if field not in target_src:
                continue
            dtype = schema.FIELD_TYPES[field]
            casted, changed = [], 0
            for row, v in out[field].items():
                new, note = self._cast(v, dtype, field)
                if note:
                    audit.append(AuditEntry(source_column=target_src[field], target_column=field, row=int(row),
                                            transformation_applied="cast_failed_blank", before_value=to_display(v),
                                            after_value=None, approved_by=self._approver(state, field),
                                            rationale=f"{note}; reviewer acknowledged the flagged issue, value left blank "
                                                      f"(no data invented)."))
                casted.append(new)
                changed += 0 if is_blank(new) else 1
            out[field] = casted
            audit.append(AuditEntry(source_column=target_src[field], target_column=field,
                                    transformation_applied=f"cast_to_{dtype}", before_value=None,
                                    after_value=f"{changed} values", approved_by="system (data dictionary)",
                                    rationale=f"Data dictionary requires {field} as {dtype}."))
        out = out.reset_index(drop=True)
        out = out.astype({f: ("Int64" if schema.FIELD_TYPES[f] == "int" else
                              "Float64" if schema.FIELD_TYPES[f] == "float" else "object")
                          for f in schema.TARGET_FIELDS})

        state.output_df, state.audit_log = out, audit
        state.summary.update({
            "processed_at": now_iso(), "rows_in": int(len(src)), "rows_out": int(len(out)),
            "rows_excluded": sorted(int(r) for r in drop_rows),
            "cells_changed": sum(1 for a in audit if a.row is not None and a.transformation_applied not in
                                 ("exclude_row",) and not a.transformation_applied.startswith("skipped")),
            "fields_blank": [f for f in schema.TARGET_FIELDS if f not in target_src],
            "recommendations": {s: sum(1 for r in state.active_recs() if r.status == s)
                                for s in ("approved", "edited", "rejected")},
        })
        problems = validate_frame(out)
        state.finish("transform", NAME, not problems,
                     "Output conforms to schema" if not problems else "; ".join(problems))
        return state

    @staticmethod
    def _approver(state: SOVState, field: str) -> str:
        for r in state.active_recs():
            if r.field == field and r.status in APPLIED and r.decided_by:
                return r.decided_by
        return "reviewer"

    @staticmethod
    def _cast(v: Any, dtype: str, field: str):
        if is_blank(v):
            return None, None
        if dtype == "string":
            s = str(to_display(v)).strip() if not isinstance(v, str) else v
            if field == "Fire Sprinklers (Y/N)" and s.strip() not in schema.SPRINKLER_CODES:
                return None, f"'{s}' is not an allowed sprinkler code (Y, N, Y13, Y(13R))"
            return s, None
        x, kind = parse_number(v)
        if x is None or kind != "ok":
            return None, f"'{v}' is not a valid {dtype}"
        if dtype == "int":
            if not float(x).is_integer():
                return None, f"'{v}' is not a whole number"
            return int(x), None
        return float(x), None


# ---------------------------------------------------------------------- validation & export
def validate_frame(df: pd.DataFrame) -> list[str]:
    problems = []
    if list(df.columns) != schema.TARGET_FIELDS:
        problems.append("Columns do not match the 17 target fields in order")
    for f in schema.TARGET_FIELDS:
        if f not in df.columns:
            continue
        t = schema.FIELD_TYPES[f]
        for v in df[f].dropna().tolist():
            if t == "string":
                ok = isinstance(v, str)
            elif t == "int":
                ok = isinstance(v, numbers.Integral) and not isinstance(v, bool)
            else:
                ok = isinstance(v, numbers.Real) and not isinstance(v, bool)
            if not ok:
                problems.append(f"{f}: value {v!r} is not {t}")
                break
    return problems


def _cell(v):
    if v is None or v is pd.NA or (isinstance(v, float) and pd.isna(v)):
        return None
    if hasattr(v, "item"):  # numpy scalar
        return v.item()
    return v


def export_files(state: SOVState) -> dict[str, bytes]:
    """Cleaned_SOV.xlsx (single sheet, no merges, no formatting), Audit_Log.xlsx, Audit_Log.json, summary."""
    df = state.output_df
    wb = Workbook()
    ws = wb.active
    ws.title = "Cleaned_SOV"
    ws.append(schema.TARGET_FIELDS)
    for row in df.itertuples(index=False):
        ws.append([_cell(v) for v in row])
    buf = io.BytesIO()
    wb.save(buf)

    audit_rows = [a.model_dump() for a in state.audit_log]
    cols = ["source_column", "target_column", "transformation_applied", "row", "before_value", "after_value",
            "confidence", "approved_by", "timestamp", "rationale", "recommendation_id"]
    wb2 = Workbook()
    ws2 = wb2.active
    ws2.title = "Audit_Log"
    ws2.append(cols)
    for a in audit_rows:
        ws2.append([_cell(a.get(c)) if not isinstance(a.get(c), (list, dict)) else json.dumps(a.get(c)) for c in cols])
    buf2 = io.BytesIO()
    wb2.save(buf2)

    summary = {
        "run_id": state.run_id, "file": state.file_name, "sheet": state.selected_sheet,
        "header_row": state.header_row, "dq_score_at_intake": state.quality.get("dq_score"),
        "mapping": state.mapping_json, **{k: v for k, v in state.summary.items() if k != "explained_ids"},
        "events": [e.model_dump() for e in state.events],
    }
    return {
        "Cleaned_SOV.xlsx": buf.getvalue(),
        "Audit_Log.xlsx": buf2.getvalue(),
        "Audit_Log.json": json.dumps(audit_rows, indent=2, default=str).encode(),
        "Processing_Summary.json": json.dumps(summary, indent=2, default=str).encode(),
    }


def validate_xlsx(data: bytes) -> list[str]:
    """Independent check of the written file (NFR-5)."""
    wb = load_workbook(io.BytesIO(data))
    problems = []
    if wb.sheetnames != ["Cleaned_SOV"]:
        problems.append(f"Sheets are {wb.sheetnames}, expected ['Cleaned_SOV']")
    ws = wb["Cleaned_SOV"] if "Cleaned_SOV" in wb.sheetnames else wb.active
    header = [c.value for c in ws[1]]
    if header != schema.TARGET_FIELDS:
        problems.append("Header row does not match target schema")
    if ws.merged_cells.ranges:
        problems.append("Merged cells present")
    for j, f in enumerate(schema.TARGET_FIELDS, start=1):
        t = schema.FIELD_TYPES[f]
        for (v,) in ws.iter_rows(min_row=2, min_col=j, max_col=j, values_only=True):
            if v is None:
                continue
            ok = isinstance(v, str) if t == "string" else (isinstance(v, int) or (isinstance(v, float) and v.is_integer())
                                                            if t == "int" else isinstance(v, (int, float)))
            if not ok:
                problems.append(f"{f}: {v!r} is not {t}")
                break
    return problems
